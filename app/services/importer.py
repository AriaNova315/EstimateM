import io
import json
from datetime import datetime, date

import openpyxl
import aiosqlite

from app.services.task_classifier import classify_tasks


def _safe(row, idx, default=None):
    try:
        v = row[idx]
        return v if v is not None else default
    except (IndexError, TypeError):
        return default


def _normalize_wbs(val) -> str | None:
    if val is None:
        return None
    if isinstance(val, int):
        return str(val)
    if isinstance(val, float):
        return str(int(val)) if val == int(val) else str(val)
    s = str(val).strip()
    return s or None


def _to_date_str(val) -> str | None:
    if val is None:
        return None
    if isinstance(val, datetime):
        return val.date().isoformat()
    if isinstance(val, date):
        return val.isoformat()
    return str(val)


def peek_project_code(content: bytes) -> str | None:
    """仅读取 Excel 中的项目编号，不做完整解析。"""
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
        ws = wb["项目信息"]
        rows = list(ws.iter_rows(min_row=2, max_row=2, values_only=True))
        r = rows[0] if rows else ()
        val = _safe(r, 1)
        return str(val).strip() if val is not None else None
    except Exception:
        return None


def _parse_excel(content: bytes) -> dict:
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)

    # ── 项目信息 ──────────────────────────────────────────
    ws = wb["项目信息"]
    rows = list(ws.iter_rows(min_row=2, max_row=2, values_only=True))
    r = rows[0] if rows else ()

    doc_est = [_safe(r, 6), _safe(r, 7), _safe(r, 8), _safe(r, 9)]
    code_est = [_safe(r, 10), _safe(r, 11), _safe(r, 12), _safe(r, 13)]
    has_estimators = any(x for x in doc_est + code_est)

    project = {
        "name": _safe(r, 0, ""),
        "code": str(_safe(r, 1, "")),
        "leader": _safe(r, 2),
        "cm_manager": _safe(r, 3),
        "qa_manager": _safe(r, 4),
        "round": int(_safe(r, 5, 1) or 1),
        "estimators": json.dumps({
            "doc": doc_est,
            "code": code_est,
            "cm": _safe(r, 14),
            "qa": _safe(r, 15),
            "mgmt": _safe(r, 16),
            "review": _safe(r, 17),
        }) if has_estimators else None,
        "team_members": _safe(r, 18),
    }

    # ── 项目进度计划 ──────────────────────────────────────
    ws = wb["项目进度计划"]
    wbs_tasks = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        name = _safe(row, 1)
        if not name:
            continue
        wbs_tasks.append({
            "wbs": _normalize_wbs(_safe(row, 0)),
            "name": str(name).strip(),
            "start_date": _to_date_str(_safe(row, 2)),
            "end_date": _to_date_str(_safe(row, 3)),
            "related_items": _safe(row, 4),
            "duration_days": _safe(row, 5),
            "workload": _safe(row, 6),
            "resources": _safe(row, 7),
        })

    # ── 配置项清单 ────────────────────────────────────────
    ws = wb["配置项清单"]
    config_items = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        name = _safe(row, 1)
        if not name:
            continue
        seq = _safe(row, 0)
        raw_phase = _safe(row, 6)
        config_items.append({
            "seq": int(seq) if seq is not None else None,
            "name": str(name).strip(),
            "identifier": _safe(row, 2),
            "version": _safe(row, 3),
            "item_type": _safe(row, 4),
            "actual_pages": _safe(row, 5),
            "phase_wbs": str(int(raw_phase)) if isinstance(raw_phase, float) and raw_phase.is_integer()
                         else (str(raw_phase).strip() if raw_phase is not None else None),
        })

    # ── 代码模块（可能不存在）────────────────────────────
    code_modules = []
    if "代码模块" in wb.sheetnames:
        ws = wb["代码模块"]
        for row in ws.iter_rows(min_row=2, values_only=True):
            loc = _safe(row, 2)
            if loc is None:
                continue
            seq = _safe(row, 0)
            name = _safe(row, 1)
            code_modules.append({
                "seq": int(seq) if seq is not None else None,
                "name": str(name).strip() if name else None,
                "actual_loc": int(loc),
            })

    return {
        "project": project,
        "wbs_tasks": wbs_tasks,
        "config_items": config_items,
        "code_modules": code_modules,
    }


async def import_excel(db: aiosqlite.Connection, project_id: int, content: bytes) -> dict:
    data = _parse_excel(content)
    p = data["project"]

    # Update project info from Excel
    await db.execute(
        """UPDATE projects SET
           name=?, code=?, leader=?, cm_manager=?, qa_manager=?,
           round=?, estimators=?, team_members=?, updated_at=CURRENT_TIMESTAMP
           WHERE id=?""",
        (p["name"], p["code"], p["leader"], p["cm_manager"], p["qa_manager"],
         p["round"], p["estimators"], p["team_members"], project_id),
    )

    # Full overwrite of related data
    # 先清除 Session 下的明细表（持有对 wbs_tasks 的 FK），再删 wbs_tasks
    await db.execute(
        "DELETE FROM sheet1_estimates WHERE session_id IN "
        "(SELECT id FROM estimate_sessions WHERE project_id=?)", (project_id,))
    await db.execute(
        "DELETE FROM code_scale_estimates WHERE session_id IN "
        "(SELECT id FROM estimate_sessions WHERE project_id=?)", (project_id,))
    await db.execute(
        "DELETE FROM activity_workload_snapshots WHERE session_id IN "
        "(SELECT id FROM estimate_sessions WHERE project_id=?)", (project_id,))
    await db.execute(
        "DELETE FROM review_workload_snapshots WHERE session_id IN "
        "(SELECT id FROM estimate_sessions WHERE project_id=?)", (project_id,))
    await db.execute("DELETE FROM estimate_sessions WHERE project_id=?", (project_id,))
    await db.execute("DELETE FROM wbs_tasks WHERE project_id=?", (project_id,))
    await db.execute("DELETE FROM config_items WHERE project_id=?", (project_id,))
    await db.execute("DELETE FROM code_modules WHERE project_id=?", (project_id,))

    # Config items
    for item in data["config_items"]:
        await db.execute(
            "INSERT INTO config_items (project_id, seq, name, identifier, version, item_type, actual_pages, phase_wbs) VALUES (?,?,?,?,?,?,?,?)",
            (project_id, item["seq"], item["name"], item["identifier"],
             item["version"], item["item_type"], item["actual_pages"], item.get("phase_wbs")),
        )

    # Code modules
    for mod in data["code_modules"]:
        await db.execute(
            "INSERT INTO code_modules (project_id, seq, name, actual_loc) VALUES (?,?,?,?)",
            (project_id, mod["seq"], mod["name"], mod["actual_loc"]),
        )

    # WBS tasks + classify
    inserted = []
    for task in data["wbs_tasks"]:
        async with db.execute(
            """INSERT INTO wbs_tasks
               (project_id, wbs, name, start_date, end_date, related_items, duration_days, workload, resources)
               VALUES (?,?,?,?,?,?,?,?,?)""",
            (project_id, task["wbs"], task["name"], task["start_date"], task["end_date"],
             task["related_items"], task["duration_days"], task["workload"], task["resources"]),
        ) as cur:
            inserted.append({"id": cur.lastrowid, "wbs": task["wbs"], "name": task["name"]})

    type_map = classify_tasks(inserted)
    for task_id, task_type in type_map.items():
        await db.execute("UPDATE wbs_tasks SET task_type=? WHERE id=?", (task_type, task_id))

    await db.commit()

    return {
        "project_id": project_id,
        "task_count": len(inserted),
        "config_item_count": len(data["config_items"]),
        "code_module_count": len(data["code_modules"]),
    }
