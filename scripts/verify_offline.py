"""Isolated HTTP regression check; never starts the application in the source directory.

Run with an interpreter containing requirements.txt (no extra test dependencies).
Temporary business data is removed when verification finishes.
"""

import hashlib
import io
import json
import os
from pathlib import Path
import shutil
import socket
import sqlite3
import subprocess
import sys
import tempfile
import time
from urllib.error import HTTPError, URLError
from urllib.request import build_opener, ProxyHandler, Request
import zipfile

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
OPENER = build_opener(ProxyHandler({}))


def main():
    original_db = ROOT / "estimatem.db"
    original_hash = hashlib.sha256(original_db.read_bytes()).digest() if original_db.exists() else None
    with tempfile.TemporaryDirectory(prefix="estimatem-offline-") as temporary:
        target = Path(temporary) / "中文 path with spaces"
        target.mkdir()
        shutil.copytree(ROOT / "app", target / "app", ignore=shutil.ignore_patterns("__pycache__"))
        with socket.socket() as sock:
            sock.bind(("127.0.0.1", 0))
            port = sock.getsockname()[1]
        base = f"http://127.0.0.1:{port}"

        def request(path, method="GET", payload=None, headers=None):
            if isinstance(payload, dict):
                payload = json.dumps(payload).encode("utf-8")
                headers = {"Content-Type": "application/json"}
            with OPENER.open(Request(base + path, data=payload, method=method, headers=headers or {}), timeout=30) as response:
                data = response.read()
                return json.loads(data) if data and "application/json" in response.headers.get("Content-Type", "") else data

        def launch():
            env = dict(os.environ, PYTHONUTF8="1", PYTHONDONTWRITEBYTECODE="1")
            env.pop("PYTHONPATH", None)
            process = subprocess.Popen(
                [sys.executable, "-m", "uvicorn", "app.main:app", "--host", "127.0.0.1", "--port", str(port)],
                cwd=target, env=env, stdout=log, stderr=log,
            )
            deadline = time.monotonic() + 20
            while time.monotonic() < deadline:
                if process.poll() is not None:
                    raise RuntimeError("Temporary server exited; see temporary log")
                try:
                    if request("/health") == {"status": "ok"}:
                        return process
                except (URLError, OSError):
                    time.sleep(0.1)
            process.terminate()
            process.wait(timeout=10)
            raise RuntimeError("Temporary server did not start")

        def stop(process):
            process.terminate()
            try:
                process.wait(timeout=10)
            except subprocess.TimeoutExpired:
                process.kill()
                process.wait(timeout=5)

        with (target / "server.log").open("wb") as log:
            process = launch()
            try:
                assert request("/api/projects") == [], "Fresh database is not empty"
                with sqlite3.connect(target / "estimatem.db") as db:
                    tables = db.execute("SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%'").fetchall()
                    assert len(tables) == 10, tables
                    assert "phase_wbs" in [r[1] for r in db.execute("PRAGMA table_info(config_items)")]
                page = request("/").decode("utf-8")
                assert 'src="/static/vendor/alpinejs-3.14.1.min.js"' in page
                assert "https://" not in page and "http://" not in page
                assert request("/static/vendor/alpinejs-3.14.1.min.js") == (ROOT / "app/static/vendor/alpinejs-3.14.1.min.js").read_bytes()
                for path in ("/docs", "/redoc"):
                    try:
                        request(path)
                        raise AssertionError(f"{path} should be disabled")
                    except HTTPError as error:
                        assert error.code == 404
                print("PASS fresh database / local JS / disabled docs / Chinese and space path")

                boundary = "EstimateMOfflineRegressionBoundary"
                content = (ROOT / "项目相关信息表.xlsx").read_bytes()
                body = (
                    f'--{boundary}\r\nContent-Disposition: form-data; name="file"; filename="sample.xlsx"\r\n'
                    'Content-Type: application/vnd.openxmlformats-officedocument.spreadsheetml.sheet\r\n\r\n'
                ).encode() + content + f"\r\n--{boundary}--\r\n".encode()
                upload_headers = {"Content-Type": f"multipart/form-data; boundary={boundary}"}
                imported = request("/api/projects/import-smart", "POST", body, upload_headers)
                project_id = imported["project_id"]
                assert imported["task_count"] > 0 and imported["config_item_count"] > 0
                tasks = request(f"/api/projects/{project_id}/tasks")
                assert len(tasks) == imported["task_count"]
                session = request(f"/api/projects/{project_id}/sessions", "POST")
                # The generator returns session_id, while list endpoints return id.
                session_id = session["session_id"]
                detail = request(f"/api/sessions/{session_id}")
                row = detail["sheet1_estimates"][0]
                value = (row["v1"] or 0) + 1
                updated = request(f'/api/sheet1-estimates/{row["id"]}', "PUT", {"v1": value})
                assert updated["v1"] == value
                output = request(f"/api/sessions/{session_id}/export")
                with zipfile.ZipFile(io.BytesIO(output)) as archive:
                    assert archive.testzip() is None
                workbook = load_workbook(io.BytesIO(output), data_only=False)
                try:
                    assert len(workbook.sheetnames) == 10
                    assert any(c.data_type == "f" for sheet in workbook for row in sheet for c in row)
                finally:
                    workbook.close()
                print("PASS sample import / task classification / session / edit / 10-sheet export with formulas")
                request(f"/api/sessions/{session_id}", "DELETE")
                assert request(f"/api/projects/{project_id}/sessions") == []
                request(f"/api/projects/{project_id}/sessions", "POST")
                imported_again = request("/api/projects/import-smart", "POST", body, upload_headers)
                assert imported_again["project_id"] == project_id and not imported_again["created"]
                assert request(f"/api/projects/{project_id}/sessions") == []
                print("PASS session deletion / same-code reimport and session reset")
            finally:
                stop(process)
            process = launch()
            try:
                assert len(request("/api/projects")) == 1
                print("PASS stop / restart / persisted temporary data")
            finally:
                stop(process)
    current_hash = hashlib.sha256(original_db.read_bytes()).digest() if original_db.exists() else None
    assert original_hash == current_hash, "Source database changed"
    print("PASS source database unchanged; Windows and Office acceptance still required")


if __name__ == "__main__":
    main()
